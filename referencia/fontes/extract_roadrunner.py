"""Adaptador do padrão RoadRunner (planilha de unifilar de restauração).
Uso: python3 extract_roadrunner.py arquivo.xlsx saida.json
Layout por linha de 20 m. Índices de coluna em base 0.
"""
import openpyxl, json, sys, re

def num(v, n=2):
    return round(v, n) if isinstance(v, (int, float)) else None
def s(v):
    return v if isinstance(v, str) and v != '' else None

def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sol_rows = list(wb['Solução'].iter_rows(values_only=True))
    U = [r for r in wb['Unifilar'].iter_rows(min_row=11, values_only=True) if isinstance(r[13], (int, float))]
    S = [r for r in sol_rows[9:] if isinstance(r[2], (int, float))]
    sol = {round(r[2] * 1000): r for r in S}
    name = sol_rows[1][2]
    rows = []
    for r in U:
        sr = sol.get(round(r[13] * 1000))
        widths = [r[176 + i] or 0 for i in range(6)]
        nl = sum(1 for w in widths if w > 0)
        lanes = []
        for n in range(nl):
            b = 21 + 17 * n
            t = 30 + 17 * n
            treat = []
            for j in range(7):
                v = r[t + j]
                treat.append(s(v) if isinstance(v, str) and v != '' else (v if v not in ('', None, 0) else 0))
            fs_tipo = r[152 + n] if r[152 + n] in ('C', 'D') else None
            esp = None; pct = None
            if sr:
                for ce, cp in ((55, 62), (69, 76)):
                    if isinstance(sr[ce + n], (int, float)) and sr[ce + n]:
                        esp = sr[ce + n]; pct = sr[cp + n]
            lanes.append([
                s(sr[7 + n]) if sr else None,
                num(r[b]), num(r[b + 1], 3), num(r[b + 2]), num(r[b + 6]), num(r[b + 7], 0), num(r[b + 8]),
                s(r[b + 3]), s(r[b + 4]), num(r[b + 5], 1), treat,
                fs_tipo, esp, pct,
                num(sr[118 + n]) if sr and isinstance(sr[118 + n], (int, float)) and sr[118 + n] else None,
            ])
        dr = None
        if sr:
            c = (sr[15], sr[18]); d = (sr[16], sr[19])
            dr = [s(c[1]) or ('?' if c[0] not in (None, '') else None),
                  s(d[1]) or ('?' if d[0] not in (None, '') else None)]
            if dr == [None, None]:
                dr = None
        rows.append([
            r[13], round(r[15], 7), round(r[16], 7), widths[:nl], r[182] or 0,
            s(sr[13]) if sr else None, s(r[19]), s(sr[21]) if sr else None, lanes,
            s(r[14]), dr, [num(r[186], 0), num(r[187], 0)],
        ])
    parts = [p.strip() for p in name.split(' - ')]
    meta = {
        'name': name, 'uf': parts[0], 'via': ' - '.join(parts[1:]),
        'travel': -1 if 'Decrescente' in name else 1,
        'kmi': sol_rows[2][2], 'kmf': sol_rows[3][2],
        'format': 'roadrunner',
    }
    return {'meta': meta, 'rows': rows}

if __name__ == '__main__':
    d = parse(sys.argv[1])
    json.dump(d, open(sys.argv[2], 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    print(d['meta'], len(d['rows']))
